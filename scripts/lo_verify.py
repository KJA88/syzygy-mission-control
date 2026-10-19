#!/usr/bin/env python3
"""Cross-check cached formula values against an independent LibreOffice recalculation.

Strips every cached <v> from formula cells in a temp copy (forcing LibreOffice to calculate them),
converts headless, and compares to the cached values the service wrote.  Needs soffice + openpyxl.
  python3 scripts/lo_verify.py workbook.xlsx [outdir]
"""
import re, shutil, subprocess, sys, tempfile, zipfile
from pathlib import Path
import openpyxl

src = Path(sys.argv[1]); out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(tempfile.mkdtemp())
out.mkdir(parents=True, exist_ok=True)
stripped = out / "stripped.xlsx"
with zipfile.ZipFile(src) as zin, zipfile.ZipFile(stripped, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename.startswith("xl/worksheets/"):
            text = data.decode("utf-8")
            text = re.sub(r"(<c\b[^>]*>)(<f>.*?</f>)<v>.*?</v>", r"\1\2", text, flags=re.S)
            text = re.sub(r'(<c\b[^>]*?) t="str"([^>]*>)(<f>)', r"\1\2\3", text)
            data = text.encode("utf-8")
        zout.writestr(item, data)
conv = out / "lo"
conv.mkdir(exist_ok=True)
subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{out}/profile", "--convert-to", "xlsx", "--outdir", str(conv), str(stripped)],
               check=True, capture_output=True, timeout=180)
lo = openpyxl.load_workbook(conv / "stripped.xlsx", data_only=True)
mine_f = openpyxl.load_workbook(src)              # formulas
mine_v = openpyxl.load_workbook(src, data_only=True)  # service-cached values
checked = mismatches = 0
errors = []
for ws in mine_f:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                a = mine_v[ws.title][c.coordinate].value
                b = lo[ws.title][c.coordinate].value
                a = None if a == "" else a
                b = None if b == "" else b
                checked += 1
                same = (a == b) or (isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) < 1e-6)
                if isinstance(b, str) and b.startswith("#"):
                    errors.append((ws.title, c.coordinate, b))
                if not same:
                    mismatches += 1
                    if mismatches <= 15:
                        print("MISMATCH", ws.title, c.coordinate, "service-cache=", a, "LibreOffice=", b)
print(f"formula cells checked: {checked}; mismatches vs LibreOffice: {mismatches}; LO error values: {len(errors)}")
sys.exit(0 if mismatches == 0 and not errors else 1)
